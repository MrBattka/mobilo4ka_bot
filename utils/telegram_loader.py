import json
import os
import asyncio
import traceback
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Dict, Any, List

import gspread
from google.oauth2.service_account import Credentials

from utils.excel_writer import save_supplier_to_excel
from utils.parsers import clean_sunrise_text
from normalizers_site.normalize_supplier_rows import normalize_supplier_rows
from config.settings import CREDENTIALS_FILE, PHONE_1, API_HASH, API_ID

TEXT_BASED_SUPPLIERS = {
    "-1001173041677",  # MiOpt
    "-1001449349991",  # MiHonor
    "-1001282286698",  # Рацмаг
    "-1002080508044",  # Arti
}

MIOPT_SUPPLIERS = {"-1001173041677"}
MIHONOR_SUPPLIERS = {"-1001449349991"}
RACMAG_SUPPLIERS = {"-1001282286698"}

try:
    from zoneinfo import ZoneInfo
except ImportError:
    try:
        import pytz
        MY_TIMEZONE = pytz.timezone('Europe/Moscow')
    except ImportError:
        class FixedMoscowTZ(timezone):
            def __init__(self):
                super().__init__(timedelta(hours=3))
        MY_TIMEZONE = FixedMoscowTZ()
        ZoneInfo = None
else:
    MY_TIMEZONE = ZoneInfo('Europe/Moscow')


def normalize_word_list(value: Any) -> List[str]:
    if value is None:
        return []

    if isinstance(value, str):
        items = [value]
    elif isinstance(value, list):
        items = value
    else:
        items = [str(value)]

    result = []
    for item in items:
        if item is None:
            continue
        s = str(item).strip().lower()
        if s:
            result.append(s)
    return result


def text_has_blocked_words(text: Any, blocked_words: List[str]) -> bool:
    if not blocked_words or text is None:
        return False

    text_value = str(text).lower()
    return any(word in text_value for word in blocked_words)


def normalize_chat_id(chat_id: Any) -> Any:
    if chat_id is None:
        return chat_id

    value = str(chat_id).strip()
    if value.startswith("@"):
        return value

    try:
        return int(value)
    except (TypeError, ValueError):
        return chat_id


def normalize_title_row(conf: Dict[str, Any]) -> List[str]:
    raw_title = conf.get("title")
    if raw_title is None:
        return []

    if isinstance(raw_title, str):
        items = [raw_title]
    elif isinstance(raw_title, (list, tuple)):
        items = raw_title
    else:
        items = [str(raw_title)]

    result = []
    for item in items:
        if item is None:
            continue
        text = str(item).strip()
        if text:
            result.append(text)

    if len(result) > 2:
        result = result[:2]

    return result


def apply_title_row(ws, title_row: List[str]):
    if not title_row:
        return

    existing = ws.get_all_values()
    if existing and existing[0][:len(title_row)] == title_row:
        return

    if len(title_row) == 1:
        ws["A1"] = title_row[0]
        return

    ws["A1"] = title_row[0]
    ws["B1"] = title_row[1]


def prepare_rows_with_title(title_row: List[str], rows: List[List[Any]]) -> List[List[Any]]:
    if not rows and not title_row:
        return []

    if not title_row:
        return rows

    target_width = max(len(title_row), max((len(row) for row in rows), default=0))

    padded_title = list(title_row) + [""] * max(0, target_width - len(title_row))
    padded_rows = []

    for row in rows:
        padded_row = list(row) + [""] * max(0, target_width - len(row))
        padded_rows.append(padded_row)

    return [padded_title] + padded_rows


def load_suppliers_config(suppliers_file: Path) -> Dict[str, Any]:
    """
    Загружает конфигурацию поставщиков из указанного файла.

    :param suppliers_file: Путь к JSON-файлу конфигурации (обязательный параметр)
    """
    try:
        with open(suppliers_file, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception as e:
        print(f"⚠️ Ошибка чтения конфига {suppliers_file}: {e}")
        return {}


def get_telethon_client():
    try:
        from telethon import TelegramClient
    except ImportError:
        print("❌ Telethon не установлен")
        return None

    if not all([API_ID, API_HASH, PHONE_1]):
        print("❌ Отсутствуют настройки Telegram")
        return None

    return TelegramClient('session_mob_bot', int(API_ID), API_HASH)


async def resolve_telethon_entity(client, chat_id: str):
    normalized = normalize_chat_id(chat_id)

    try:
        entity = await client.get_entity(normalized)
        if entity:
            return entity
    except Exception as e:
        print(f"⚠️ Стандартный get_entity не сработал для {chat_id}: {e}")

    try:
        from telethon.tl.types import PeerUser, PeerChannel, PeerChat
    except ImportError:
        print("❌ Telethon не установлен")
        return None

    if isinstance(normalized, int):
        for peer_type in (PeerChannel, PeerChat, PeerUser):
            try:
                entity = await client.get_entity(peer_type(normalized))
                if entity:
                    return entity
            except Exception as e:
                print(f"⚠️ Попытка {peer_type.__name__}({normalized}) не удалась: {e}")

    if isinstance(normalized, str):
        if normalized.startswith("-100"):
            try:
                channel_id = int(normalized[4:])
                entity = await client.get_entity(PeerChannel(channel_id))
                if entity:
                    return entity
            except Exception as e:
                print(f"⚠️ Попытка PeerChannel для {normalized}: {e}")
        elif normalized.startswith("-"):
            try:
                chat_id_int = int(normalized[1:])
                entity = await client.get_entity(PeerChat(chat_id_int))
                if entity:
                    return entity
            except Exception as e:
                print(f"⚠️ Попытка PeerChat для {normalized}: {e}")

    return None


async def sync_suppliers_to_sheets(
    config: Dict[str, Any],
    sheet_id: str,
    suppliers_file: Path,
    use_excel_mode: bool = False
):
    """
    Синхронизирует поставщиков.
    Если use_excel_mode=True — выгружает в Excel, иначе в Google Sheets.
    """
    print(f"📂 Синхронизация начата для конфига: {suppliers_file}")

    client = get_telethon_client()
    if not client:
        print(f"❌ Не удалось создать клиент Telegram для {suppliers_file}")
        return False

    await client.start(phone=os.getenv('PHONE_1') or PHONE_1)
    
    try:
        if use_excel_mode:
            for chat_id, conf in config.items():
                if conf.get("skip", False):
                    print(
                        f"⏭️ Пропускаем поставщика: "
                        f"{conf.get('name', chat_id)} (skip=True)"
                    )
                    continue

                await process_supplier(
                    client,
                    chat_id,
                    conf,
                    sh=None,
                    use_excel_mode=True,
                )

            print(
                f"✅ Готово. Данные записаны в Excel "
                f"(конфиг: {suppliers_file.name})"
            )
            return True

        try:
            creds = Credentials.from_service_account_file(
                CREDENTIALS_FILE,
                scopes=[
                    "https://www.googleapis.com/auth/spreadsheets"
                ],
            )
            gc = gspread.authorize(creds)
            sh = gc.open_by_key(sheet_id)
        except Exception as e:
            print(
                f"❌ Ошибка Sheets (ID: {sheet_id}) "
                f"для {suppliers_file}: {e}"
            )
            return False

        for chat_id, conf in config.items():
            if conf.get("skip", False):
                print(
                    f"⏭️ Пропускаем поставщика: "
                    f"{conf.get('name', chat_id)} (skip=True)"
                )
                continue

            await process_supplier(
                client,
                chat_id,
                conf,
                sh=sh,
                use_excel_mode=False,
            )

        print(
            f"✅ Готово. Данные записаны в таблицу {sheet_id} "
            f"(конфиг: {suppliers_file.name})"
        )
        return True

    except Exception as e:
        print(f"❌ Ошибка при синхронизации для {suppliers_file}: {e}")
        traceback.print_exc()
        return False

    finally:
        await client.disconnect()


async def process_supplier(client, chat_id: str, conf: Dict[str, Any], sh=None, use_excel_mode: bool = False):
    if conf.get("skip", False):
        print(f"⏭️ Пропускаем поставщика: {conf.get('name', chat_id)}")
        return
    
    if conf.get("isLocal", False):
        print(f"📁 Локальный источник: {conf.get('name', chat_id)} — пропускаем Telegram-поиск и читаем уже записанный лист")
        return

    name = conf.get("name", "Unknown").lower()
    sheet_number = conf.get("numberList", 1)
    replace_old = conf.get("replaceOld", True)
    message_ids = conf.get("messageIds", None)
    today_only = conf.get("todayMessage", False)
    title_row = normalize_title_row(conf)

    is_miopts = chat_id in MIOPT_SUPPLIERS
    is_mihonor = chat_id in MIHONOR_SUPPLIERS
    is_racmag = chat_id in RACMAG_SUPPLIERS

    trigger_cmd = conf.get("txtMessageHandle", "").strip()
    is_txt_handle = bool(trigger_cmd)

    read_excel = conf.get("readExcel", False)
    excel_columns = conf.get("excelColumns", {})

    raw_num_messages = conf.get("numMessages")
    if isinstance(raw_num_messages, str):
        try:
            raw_num_messages = int(raw_num_messages)
        except ValueError:
            raw_num_messages = None

    message_limit = raw_num_messages if isinstance(raw_num_messages, int) and raw_num_messages > 0 else 100
    parse_all_messages = bool(conf.get("parseAllMessages", False))
    effective_limit = None if parse_all_messages else message_limit

    message_code_words = normalize_word_list(conf.get("messageCodeWord", ""))
    ban_code_words = normalize_word_list(conf.get("banCodeWord", []))

    has_keyword = bool(message_code_words)
    keyword = ", ".join(message_code_words) if message_code_words else ""

    has_specific_ids = message_ids is not None and len(message_ids) > 0

    print(
        f"🔍 Обработка: {name} ({chat_id}), handle: {trigger_cmd}, "
        f"read_excel: {read_excel}, keyword: {keyword}, banned: {ban_code_words}, "
        f"numMessages: {message_limit}, parseAllMessages: {parse_all_messages}, "
        f"title: {title_row}"
    )

    normalized_chat_id = normalize_chat_id(chat_id)
    try:
        chat_id_int = int(normalized_chat_id) if isinstance(normalized_chat_id, int) else None
        entity_arg = normalized_chat_id
        is_numeric_id = isinstance(normalized_chat_id, int)
    except Exception:
        chat_id_int = None
        entity_arg = chat_id
        is_numeric_id = False

    if not use_excel_mode:
        try:
            if sh is None:
                raise ValueError("Google Sheets client not provided")
            if len(sh.worksheets()) >= sheet_number:
                ws = sh.worksheets()[sheet_number - 1]
            else:
                print(f"⚠️ Лист {sheet_number} не найден, создаем...")
                ws = sh.add_worksheet(title=f"Supplier_{sheet_number}", rows=1000, cols=10)
        except Exception as e:
            print(f"❌ Ошибка листа: {e}")
            return

        if replace_old:
            ws.clear()

    rows_to_write = []

    if is_txt_handle:
        print(f"📩 Отправка запроса бота: '{trigger_cmd}'...")
        try:
            entity_for_send = await resolve_telethon_entity(client, chat_id)

            if entity_for_send is None:
                print(f"❌ Ошибка разрешения сущности чата {chat_id}: Не удалось найти объект.")
                return

            try:
                sent_messages = await client.send_message(entity_for_send, trigger_cmd)
            except Exception as send_err:
                print(f"❌ Не удалось отправить сообщение в сущность {chat_id}: {send_err}")
                return

            if not sent_messages:
                print("❌ Не удалось отправить команду")
                return

            my_sent_id = sent_messages.id
            print(f"⏳ Ожидание ответа бота (10 сек)...")
            await asyncio.sleep(10)

            try:
                if is_numeric_id:
                    recent_msgs = await client.get_messages(chat_id_int, limit=effective_limit)
                else:
                    recent_msgs = await client.get_messages(entity_arg, limit=effective_limit)

                collected_messages = []
                for msg in recent_msgs:
                    if getattr(msg, "out", False):
                        continue
                    if getattr(msg, "id", 0) <= my_sent_id:
                        continue

                    msg_text = getattr(msg, "text", "")
                    if text_has_blocked_words(msg_text, ban_code_words):
                        continue

                    collected_messages.append(msg)

            except Exception as e:
                print(f"⚠️ Ошибка получения сообщений через get_messages: {e}")
                collected_messages = []

            if not collected_messages:
                print("⚠️ Новые сообщения от бота не обнаружены после фильтрации.")
            else:
                sorted_messages = sorted(collected_messages, key=lambda m: m.id)

                if read_excel:
                    print(f"📥 Режим Excel: Поиск файлов в последних {message_limit} сообщениях...")
                    excel_data_rows = []

                    for msg in sorted_messages:
                        document_obj = getattr(msg, "document", None)

                        if not document_obj and getattr(msg, "media", None):
                            media = msg.media
                            if hasattr(media, "document") and media.document:
                                document_obj = media.document

                        if not document_obj:
                            continue

                        file_name = getattr(document_obj, "file_name", None)
                        mime = getattr(document_obj, "mime_type", "") or ""

                        if not file_name:
                            for attr in getattr(document_obj, "attributes", []) or []:
                                if hasattr(attr, "file_name") and attr.file_name:
                                    file_name = attr.file_name
                                    break

                        if not file_name:
                            file_name = f"price_list_{msg.id}.xlsx"

                        is_excel = (
                            file_name.lower().endswith((".xlsx", ".xls"))
                            or mime in [
                                "application/vnd.ms-excel",
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                            ]
                        )

                        if not is_excel:
                            continue

                        file_path = str(Path("temp_files") / file_name)
                        print(f"📄 Найден Excel документ: {file_name} (MIME: {mime})")

                        try:
                            Path("temp_files").mkdir(exist_ok=True)
                            await client.download_media(msg, file=file_path)

                            parsed_rows = parse_excel_file_with_columns(file_path, excel_columns)
                            excel_data_rows.extend(parsed_rows)
                            print(f"✅ Распарсено строк из Excel ({file_name}): {len(parsed_rows)}")
                        except Exception as e:
                            print(f"❌ Ошибка обработки файла {file_name}: {e}")
                            traceback.print_exc()
                        finally:
                            if Path(file_path).exists():
                                Path(file_path).unlink()

                    if excel_data_rows:
                        rows_to_write = excel_data_rows
                        print(f"✅ Загружено {len(rows_to_write)} строк из Excel.")
                    else:
                        print("⚠️ Не удалось найти или распарсить Excel файлы.")
                        print(f"❌ Бот '{name}' в данный момент закрыт или не прислал прайс-лист.")

                else:
                    bot_text_parts = []
                    for msg in sorted_messages:
                        if getattr(msg, "text", None):
                            bot_text_parts.append(msg.text)

                    if bot_text_parts:
                        full_bot_response = "\n".join(bot_text_parts)
                        print(f"📥 Получен полный ответ от бота ({len(full_bot_response)} символов)...")

                        if "sunrise" in name:
                            cleaned_response = clean_sunrise_text(full_bot_response)
                        else:
                            cleaned_response = full_bot_response

                        lines = cleaned_response.split('\n')
                        products_count = 0
                        for line in lines:
                            clean_line = line.strip()
                            if not clean_line:
                                continue
                            if "sunrise" in name:
                                if "Актуально на" in clean_line or "🚚 Как понять" in clean_line:
                                    continue

                            rows_to_write.append([clean_line, ""])
                            products_count += 1

                        print(f"✅ Парсинг завершен. Найдено {products_count} товаров/строк.")
                    else:
                        print("⚠️ Бот прислал только медиа или пустые сообщения.")

        except Exception as e:
            print(f"❌ Ошибка при работе с ботом txtMessageHandle: {e}")
            traceback.print_exc()

    elif has_keyword:
        print(f"🔍 Поиск сообщений с ключевыми словами {message_code_words} для {name}...")
        entity = None
        try:
            entity = await client.get_entity(normalized_chat_id)
        except Exception as e:
            print(f"❌ Ошибка получения сущности чата {name}: {e}")
            return

        matched_messages_by_word = {}
        matched_messages = []

        try:
            async for msg in client.iter_messages(entity, limit=effective_limit):
                if not msg.text:
                    continue

                if text_has_blocked_words(msg.text, ban_code_words):
                    continue

                msg_text_lower = msg.text.lower()
                for word in message_code_words:
                    if word not in matched_messages_by_word and word in msg_text_lower:
                        matched_messages_by_word[word] = msg
                        if msg not in matched_messages:
                            matched_messages.append(msg)

                if len(matched_messages_by_word) == len(message_code_words):
                    break
        except Exception as e:
            print(f"❌ Ошибка поиска сообщений по ключевым словам для {name}: {e}")

        if matched_messages:
            print(f"✅ Найдено {len(matched_messages)} сообщения(й) для ключевых слов.")

            for msg in matched_messages:
                text_to_parse = msg.text
                if "sunrise" in name:
                    cleaned_text = clean_sunrise_text(text_to_parse)
                    lines = cleaned_text.split('\n')
                else:
                    lines = text_to_parse.split('\n')

                for line in lines:
                    clean_line = line.strip()
                    if not clean_line:
                        continue

                    if "sunrise" in name:
                        if "Актуально на" in clean_line or "🚚 Как понять" in clean_line:
                            continue

                    rows_to_write.append([clean_line, ""])
        else:
            print(f"⚠️ Ни одно сообщение с ключевыми словами {message_code_words} не найдено среди последних {message_limit} сообщений для {name}.")

    elif has_specific_ids:
        entity = None
        try:
            entity = await client.get_entity(normalized_chat_id)
        except Exception as e:
            print(f"❌ Ошибка чата {name}: {e}")
            return

        messages = []
        for mid in message_ids:
            try:
                msgs = await client.get_messages(entity, ids=[mid])
                if msgs:
                    if isinstance(msgs, list):
                        messages.extend(msgs)
                    elif hasattr(msgs, '__iter__') and not isinstance(msgs, (str, bytes)):
                        messages.extend(list(msgs))
                    else:
                        messages.append(msgs)
            except Exception as e:
                print(f"⚠️ Ошибка получения сообщения ID {mid} для {name}: {e}")

        all_text_parts = []
        for msg in messages:
            if not msg:
                continue

            if text_has_blocked_words(getattr(msg, "text", ""), ban_code_words):
                continue

            if getattr(msg, "text", None):
                all_text_parts.append(msg.text)
            elif getattr(msg, "media", None) and getattr(msg.media, "caption", None):
                all_text_parts.append(msg.media.caption)

        if all_text_parts:
            full_response = "\n".join(all_text_parts)
            print(f"📥 Получен полный ответ от бота ({len(full_response)} символов)...")

            if "sunrise" in name:
                cleaned_response = clean_sunrise_text(full_response)
            else:
                cleaned_response = full_response

            lines = cleaned_response.split('\n')
            products_count = 0
            for line in lines:
                clean_line = line.strip()
                if not clean_line:
                    continue

                if "sunrise" in name:
                    if "Актуально на" in clean_line or "🚚 Как понять" in clean_line:
                        continue

                rows_to_write.append([clean_line, ""])
                products_count += 1

            print(f"✅ Парсинг завершен. Найдено {products_count} товаров/строк.")
        else:
            print("⚠️ Не удалось извлечь текст из сообщений.")

    elif read_excel:
        print(f"📥 Excel-mode для {name}: проверяем последние {message_limit} сообщений...")
        entity = None
        try:
            entity = await client.get_entity(normalized_chat_id)
        except Exception as e:
            print(f"❌ Ошибка получения сущности для Excel {name}: {e}")
            return

        excel_data_rows = []
        try:
            async for msg in client.iter_messages(entity, limit=effective_limit):
                document_obj = getattr(msg, "document", None)

                if not document_obj and getattr(msg, "media", None):
                    media = msg.media
                    if hasattr(media, "document") and media.document:
                        document_obj = media.document

                if not document_obj:
                    continue

                file_name = getattr(document_obj, "file_name", None)
                mime = getattr(document_obj, "mime_type", "") or ""

                if not file_name:
                    for attr in getattr(document_obj, "attributes", []) or []:
                        if hasattr(attr, "file_name") and attr.file_name:
                            file_name = attr.file_name
                            break

                if not file_name:
                    file_name = f"price_list_{msg.id}.xlsx"

                is_excel = (
                    file_name.lower().endswith((".xlsx", ".xls"))
                    or mime in [
                        "application/vnd.ms-excel",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ]
                )

                if not is_excel:
                    continue

                file_path = str(Path("temp_files") / file_name)
                print(f"📄 Найден Excel документ: {file_name} (MIME: {mime})")

                try:
                    Path("temp_files").mkdir(exist_ok=True)
                    await client.download_media(msg, file=file_path)

                    parsed_rows = parse_excel_file_with_columns(file_path, excel_columns)
                    excel_data_rows.extend(parsed_rows)
                    print(f"✅ Excel: {file_name} -> {len(parsed_rows)} строк")
                except Exception as e:
                    print(f"❌ Ошибка обработки Excel файла {file_name}: {e}")
                    traceback.print_exc()
                finally:
                    if Path(file_path).exists():
                        Path(file_path).unlink()

        except Exception as e:
            print(f"❌ Ошибка чтения последних сообщений Excel для {name}: {e}")
            traceback.print_exc()

        if excel_data_rows:
            rows_to_write = excel_data_rows
            print(f"✅ Загружено {len(rows_to_write)} строк из Excel для {name}.")
        else:
            print(f"⚠️ Excel для {name} не найден среди последних {message_limit} сообщений.")

    elif is_miopts:
        entity = None
        try:
            entity = await client.get_entity(normalized_chat_id)
        except Exception as e:
            print(f"❌ Ошибка чата MiOpt: {e}")
            return

        messages = []
        msk_today = datetime.now(MY_TIMEZONE).date()

        if message_ids:
            for mid in message_ids:
                try:
                    msgs = await client.get_messages(entity, ids=[mid])
                    if msgs:
                        if isinstance(msgs, list):
                            messages.extend(msgs)
                        elif hasattr(msgs, '__iter__') and not isinstance(msgs, (str, bytes)):
                            messages.extend(list(msgs))
                        else:
                            messages.append(msgs)
                except Exception as e:
                    print(f"❌ Ошибка ID {mid}: {e}")
        else:
            try:
                async for msg in client.iter_messages(entity, limit=effective_limit):
                    if msg.date is None:
                        continue
                    if not msg.text:
                        continue
                    if text_has_blocked_words(msg.text, ban_code_words):
                        continue
                    if today_only:
                        msg_date_msk = msg.date.replace(tzinfo=timezone.utc).astimezone(MY_TIMEZONE)
                        if msg_date_msk.date() != msk_today:
                            continue
                    messages.append(msg)
            except Exception as e:
                print(f"❌ Ошибка загрузки сообщений MiOpt: {e}")
                traceback.print_exc()

        for msg in messages:
            if msg.text:
                lines = msg.text.split('\n')
                for line in lines:
                    clean_line = line.strip()
                    if clean_line:
                        rows_to_write.append([clean_line, ""])
        print(f"📥 MiOpt: Добавлено {len(rows_to_write)} строк.")

    elif is_mihonor:
        entity = None
        try:
            entity = await client.get_entity(normalized_chat_id)
        except Exception as e:
            print(f"❌ Ошибка чата MiHonor: {e}")
            return

        messages = []
        msk_today = datetime.now(MY_TIMEZONE).date()

        if message_ids:
            for mid in message_ids:
                try:
                    msgs = await client.get_messages(entity, ids=[mid])
                    if msgs:
                        if isinstance(msgs, list):
                            messages.extend(msgs)
                        elif hasattr(msgs, '__iter__') and not isinstance(msgs, (str, bytes)):
                            messages.extend(list(msgs))
                        else:
                            messages.append(msgs)
                except Exception as e:
                    print(f"❌ Ошибка ID {mid}: {e}")
        else:
            try:
                async for msg in client.iter_messages(entity, limit=effective_limit):
                    if msg.date is None:
                        continue
                    if not msg.text:
                        continue
                    if text_has_blocked_words(msg.text, ban_code_words):
                        continue
                    if today_only:
                        msg_date_msk = msg.date.replace(tzinfo=timezone.utc).astimezone(MY_TIMEZONE)
                        if msg_date_msk.date() != msk_today:
                            continue
                    messages.append(msg)
            except Exception as e:
                print(f"❌ Ошибка загрузки сообщений MiHonor: {e}")
                traceback.print_exc()

        for msg in messages:
            if msg.text:
                lines = msg.text.split('\n')
                for line in lines:
                    clean_line = line.strip()
                    if clean_line:
                        rows_to_write.append([clean_line, ""])
        print(f"📥 MiHonor: Добавлено {len(rows_to_write)} строк.")

    elif is_racmag:
        entity = None
        try:
            entity = await client.get_entity(normalized_chat_id)
        except Exception as e:
            print(f"❌ Ошибка чата Рацмаг: {e}")
            return

        messages = []
        msk_today = datetime.now(MY_TIMEZONE).date()

        if message_ids:
            for mid in message_ids:
                try:
                    msgs = await client.get_messages(entity, ids=[mid])
                    if msgs:
                        if isinstance(msgs, list):
                            messages.extend(msgs)
                        elif hasattr(msgs, '__iter__') and not isinstance(msgs, (str, bytes)):
                            messages.extend(list(msgs))
                        else:
                            messages.append(msgs)
                except Exception as e:
                    print(f"❌ Ошибка ID {mid}: {e}")
        else:
            try:
                async for msg in client.iter_messages(entity, limit=effective_limit):
                    if msg.date is None:
                        continue
                    if not msg.text:
                        continue
                    if text_has_blocked_words(msg.text, ban_code_words):
                        continue
                    if today_only:
                        msg_date_msk = msg.date.replace(tzinfo=timezone.utc).astimezone(MY_TIMEZONE)
                        if msg_date_msk.date() != msk_today:
                            continue
                    messages.append(msg)
            except Exception as e:
                print(f"❌ Ошибка загрузки сообщений Рацмаг: {e}")
                traceback.print_exc()

        for msg in messages:
            if msg.text:
                lines = msg.text.split('\n')
                for line in lines:
                    clean_line = line.strip()
                    if clean_line:
                        rows_to_write.append([clean_line, ""])
        print(f"📥 Рацмаг: Добавлено {len(rows_to_write)} строк.")

    else:
        print(f"🔍 Универсальный режим для {name}...")

        try:
            entity = await client.get_entity(normalized_chat_id)
        except Exception as e:
            print(f"❌ Ошибка получения сущности для {name}: {e}")
            return

        messages = []
        msk_today = datetime.now(MY_TIMEZONE).date()

        try:
            async for msg in client.iter_messages(entity, limit=effective_limit):
                if msg.date is None:
                    continue
                if not msg.text:
                    continue
                if text_has_blocked_words(msg.text, ban_code_words):
                    continue

                if today_only:
                    msg_date_msk = msg.date.replace(tzinfo=timezone.utc).astimezone(MY_TIMEZONE)
                    if msg_date_msk.date() != msk_today:
                        continue

                messages.append(msg)
        except Exception as e:
            print(f"❌ Ошибка чтения сообщений для {name}: {e}")
            traceback.print_exc()

        for msg in messages:
            if msg.text:
                lines = msg.text.split('\n')
                for line in lines:
                    clean_line = line.strip()
                    if clean_line:
                        rows_to_write.append([clean_line, ""])

        print(f"📥 {name}: Добавлено {len(rows_to_write)} строк из истории.")

    if rows_to_write:
        rows_to_write = normalize_supplier_rows(
            conf.get("normalizer", ""),
            rows_to_write,
        )

        if not rows_to_write:
            print(f"⚠️ После нормализации нет данных для {name}.")
            return
        title_row = normalize_title_row(conf)

        if title_row and replace_old:
            rows_to_write = prepare_rows_with_title(title_row, rows_to_write)
        elif title_row and not replace_old and rows_to_write and rows_to_write[0][:len(title_row)] != title_row:
            rows_to_write = [title_row] + rows_to_write

        if use_excel_mode:
            save_supplier_to_excel(
                chat_id=str(chat_id),
                name=conf.get("name", "Unknown"),
                rows=rows_to_write,
                mode="replace" if replace_old else "append",
                sheet_number=conf.get("numberList", 1)
            )
            print(f"💾 Excel export: {conf.get('name', 'Unknown')} -> {len(rows_to_write)} rows")
            return

        if replace_old:
            if sh is not None:
                ws.clear()
                if title_row:
                    apply_title_row(ws, title_row)
                data_rows = rows_to_write
                if title_row and rows_to_write and rows_to_write[0][:len(title_row)] == title_row:
                    data_rows = rows_to_write[1:]
                if data_rows:
                    ws.update("A2", data_rows if len(title_row) == 1 else data_rows)
        else:
            try:
                if title_row:
                    apply_title_row(ws, title_row)
                ws.append_rows(rows_to_write)
                print(f"📝 Записано: {len(rows_to_write)} строк для {name} (batch mode)")
            except Exception as e:
                print(f"⚠️ Ошибка batch-записи для {name}: {e}. Пробую по одной...")
                for row in rows_to_write:
                    ws.append_row(row)
    else:
        print(f"ℹ️ Пусто для {name}.")
        if title_row:
            if use_excel_mode:
                save_supplier_to_excel(
                    chat_id=str(chat_id),
                    name=conf.get("name", "Unknown"),
                    rows=[title_row],
                    mode="replace" if replace_old else "append",
                    sheet_number=conf.get("numberList", 1)
                )
            elif sh is not None and replace_old:
                ws.clear()
                apply_title_row(ws, title_row)


def parse_excel_file(file_path: str) -> List[Any]:
    try:
        lower_path = file_path.lower()
        if lower_path.endswith(".xls"):
            import xlrd
            wb = xlrd.open_workbook(file_path)
            ws = wb.sheet_by_index(0)
            data = []
            for r in range(1, ws.nrows):
                row = ws.row_values(r)
                if any(cell is not None and str(cell).strip() != "" for cell in row):
                    data.append([str(c) if c is not None else "" for c in row])
            return data

        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(cell is not None for cell in row):
                data.append([str(c) if c is not None else "" for c in row])
        return data
    except Exception as e:
        print(f"❌ Excel error: {e}")
        return []


def parse_excel_file_with_columns(file_path: str, col_config: Dict[str, int]) -> List[List[Any]]:
    """
    Парсит Excel файл и возвращает список списков, где каждый внутренний список
    содержит значения только из указанных колонок.

    col_config: dict, например {"name": 0, "price": 1}
    Возвращает: [[val_name, val_price], [val_name, val_price], ...]
    """
    try:
        lower_path = file_path.lower()

        if lower_path.endswith(".xls"):
            import xlrd
            wb = xlrd.open_workbook(file_path)
            ws = wb.sheet_by_index(0)

            result_rows = []
            columns_to_read = list(col_config.items())

            for r in range(1, ws.nrows):
                row = ws.row_values(r)
                if not any(cell is not None and str(cell).strip() != "" for cell in row):
                    continue

                extracted_row = []
                for _, col_index in columns_to_read:
                    val = row[col_index] if col_index < len(row) else ""
                    extracted_row.append(str(val).strip() if val is not None else "")
                result_rows.append(extracted_row)

            return result_rows

        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        result_rows = []
        columns_to_read = list(col_config.items())

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(cell is not None for cell in row):
                continue

            extracted_row = []
            for _, col_index in columns_to_read:
                val = row[col_index] if col_index < len(row) else None
                extracted_row.append(str(val) if val is not None else "")
            result_rows.append(extracted_row)

        return result_rows
    except Exception as e:
        print(f"❌ Excel error with columns parsing: {e}")
        traceback.print_exc()
        return []