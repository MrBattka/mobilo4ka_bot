from pathlib import Path
from typing import Any

import gspread
from google.oauth2.service_account import Credentials
from openpyxl import Workbook, load_workbook

from config.settings import (
    ALL_PRICE_FILE,
    CREDENTIALS_FILE,
    DATA_DIR_SITE,
    SELECTED_FILES,
    SHEETS_ID_FOR_SITE,
)
from defaultFixName import default_fix_name
from normalizers_site.base import fix_name_base
from returnNameWithIDForSite import match_product_for_site
from normalizers_site.ban_words import ban_word
from normalizers_site.boltun import return_fix_name_boltun
from normalizers_site.trubkoved import return_fix_name_trubkoved
from normalizers_site.applegod import return_fix_name_apple_god
from normalizers_site.store77 import return_fix_name_store77

OUTPUT_FILE_WITH_ID = Path(ALL_PRICE_FILE).with_name("AllPriceWithID.xlsx")
OUTPUT_FILE_NOT_ID = Path(ALL_PRICE_FILE).with_name("AllPriceNotID.xlsx")

GOOGLE_SHEETS_SCOPE = ["https://www.googleapis.com/auth/spreadsheets"]


def _normalize_product_name(
    name: str,
    supplier_name: str,
    is_base_supplier: bool,
) -> str:
    """Применяет нормализатор поставщика, затем общую нормализацию."""
    supplier_normalizers = {
        "boltun": return_fix_name_boltun,
        "trubkoved": return_fix_name_trubkoved,
        "applegod": return_fix_name_apple_god,
        "store77": return_fix_name_store77,
    }

    if is_base_supplier:
        normalized_name = fix_name_base(name)
    else:
        normalizer = supplier_normalizers.get(supplier_name.strip().lower())
        normalized_name = normalizer(name) if normalizer else name

    return default_fix_name(normalized_name or "")


def _normalize_sheet_value(value: Any) -> Any:
    if value is None:
        return ""

    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)

    text = str(value).strip()
    if text.startswith(("'", '"')):
        text = text[1:]

    if text == "":
        return ""

    try:
        if "." in text:
            numeric = float(text)
            return int(numeric) if numeric.is_integer() else numeric
        return int(text)
    except ValueError:
        return text


def _get_google_sheet_client():
    if not CREDENTIALS_FILE:
        raise RuntimeError("CREDENTIALS_FILE is not configured")

    creds_path = Path(CREDENTIALS_FILE)
    if not creds_path.exists():
        raise FileNotFoundError(f"Google credentials file not found: {creds_path}")

    credentials = Credentials.from_service_account_file(
        str(creds_path),
        scopes=GOOGLE_SHEETS_SCOPE,
    )
    return gspread.authorize(credentials)


def _write_google_sheets(rows_with_id: list[list[Any]]) -> None:
    if not SHEETS_ID_FOR_SITE:
        raise RuntimeError("SHEETS_ID_FOR_SITE is not configured")

    client = _get_google_sheet_client()
    spreadsheet = client.open_by_key(str(SHEETS_ID_FOR_SITE))

    sheets = spreadsheet.worksheets()
    title_to_sheet = {ws.title: ws for ws in sheets}

    ws_with_id = title_to_sheet.get("AllPriceWithID")
    if ws_with_id is None:
        ws_with_id = spreadsheet.add_worksheet(
            title="AllPriceWithID",
            rows=str(max(100, len(rows_with_id) + 10)),
            cols="20",
        )

    ws_with_id.clear()
    ws_with_id.append_rows(rows_with_id, value_input_option="RAW")


def _save_local_not_id(rows_not_id: list[list[Any]]) -> None:
    if not rows_not_id:
        OUTPUT_FILE_NOT_ID.parent.mkdir(parents=True, exist_ok=True)
        Workbook().save(OUTPUT_FILE_NOT_ID)
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "AllPriceNotID"
    ws.append(["Название", "Цена", "Поставщик"])

    for row in rows_not_id:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column, width in {"A": 60, "B": 15, "C": 25}.items():
        ws.column_dimensions[column].width = width

    OUTPUT_FILE_NOT_ID.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE_NOT_ID)

    print(f"✅ Локальный Not ID сохранён: {OUTPUT_FILE_NOT_ID}")


def _load_product_data() -> dict[str, dict[str, Any]]:
    result = {}

    for filename in SELECTED_FILES:
        file_path = Path(DATA_DIR_SITE) / filename

        if not file_path.exists():
            continue

        try:
            with file_path.open("r", encoding="utf-8") as file:
                data = __import__("json").load(file)

            if isinstance(data, dict):
                result.update(data)

        except Exception as error:
            print(f"⚠️ Ошибка загрузки базы ID {filename}: {error}")

    return result


def _parse_price(value: Any) -> int | float | None:
    if value in (None, ""):
        return None

    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()
    text = text.replace("\u00a0", "").replace(" ", "")
    text = text.replace("₽", "").replace("руб.", "").replace(",", ".")

    try:
        number = float(text)
    except ValueError:
        return None

    return int(number) if number.is_integer() else number


def _is_header(name: str, price: Any) -> bool:
    name_lower = name.strip().lower()

    if name_lower in {
        "name",
        "наименование",
        "модификация",
        "цена",
        "price",
        "hi",
        "garmin",
        "infinity",
    }:
        return True

    if str(price).strip().lower() in {
        "price",
        "цена",
        "прайс",
    }:
        return True

    return False


def export_all_price_with_id(
    suppliers_config: dict[str, dict[str, Any]] | None = None,
) -> tuple[int, int]:
    if not Path(ALL_PRICE_FILE).exists():
        print(f"⚠️ Файл не найден: {ALL_PRICE_FILE}")
        return 0, 0

    source_wb = load_workbook(ALL_PRICE_FILE, data_only=True)

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = "AllPriceWithID"
    output_ws.append(["ID", "Название", "Цена", "Поставщик"])

    product_data = _load_product_data()

    supplier_by_sheet_number: dict[int, str] = {
        25: "Boltun",
        27: "Trubkoved",
        28: "AppleGod",
        29: "Store77",
    }

    if suppliers_config:
        for conf in suppliers_config.values():
            if not isinstance(conf, dict):
                continue

            sheet_number = conf.get("numberList")
            supplier_name = conf.get("name")

            try:
                sheet_number = int(sheet_number)
            except (TypeError, ValueError):
                continue

            if supplier_name:
                supplier_by_sheet_number[sheet_number] = str(supplier_name)

    total_rows = 0
    linked_rows = 0
    not_linked_rows = 0
    rows_not_id: list[list[Any]] = []

    for sheet_index, ws in enumerate(source_wb.worksheets, start=1):
        supplier_name = supplier_by_sheet_number.get(sheet_index, ws.title)

        for row in ws.iter_rows(values_only=True):
            if not row:
                continue

            is_base_supplier = supplier_name == "База" or sheet_index == 15

            if is_base_supplier:
                raw_name = row[2] if len(row) > 2 else ""
                raw_price = row[3] if len(row) > 3 else ""
            else:
                raw_name = row[0] if len(row) > 0 else ""
                raw_price = row[1] if len(row) > 1 else ""

            name = str(raw_name or "").strip()
            price = _parse_price(raw_price)

            if not name or price is None:
                continue

            if _is_header(name, raw_price):
                continue
            
            if ban_word(name):
                continue

            total_rows += 1

            normalized_name = _normalize_product_name(
                name=name,
                supplier_name=supplier_name,
                is_base_supplier=is_base_supplier,
            )

            if not normalized_name:
                continue
            
            if ban_word(normalized_name):
                continue

            product_id = match_product_for_site(normalized_name, product_data)

            if not product_id:
                rows_not_id.append([name, price, supplier_name])
                not_linked_rows += 1
                continue

            product_id = _normalize_sheet_value(product_id)

            output_ws.append([
                product_id,
                name,
                price,
                supplier_name,
            ])
            linked_rows += 1

    output_ws.freeze_panes = "A2"
    output_ws.auto_filter.ref = output_ws.dimensions

    for column, width in {"A": 18, "B": 60, "C": 15, "D": 5, "E": 25}.items():
        output_ws.column_dimensions[column].width = width

    OUTPUT_FILE_WITH_ID.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(OUTPUT_FILE_WITH_ID)

    rows_with_id = list(output_ws.iter_rows(values_only=True))
    rows_with_id = [
        [_normalize_sheet_value(cell) for cell in row]
        for row in rows_with_id
    ]

    try:
        _write_google_sheets(rows_with_id)
        print(f"✅ With ID выгружен в Google Sheet: {SHEETS_ID_FOR_SITE}")
    except Exception as error:
        print(f"⚠️ Ошибка записи With ID в Google Sheet: {error}")

    _save_local_not_id(rows_not_id)

    print(
        f"✅ AllPriceWithID создан: {OUTPUT_FILE_WITH_ID}. "
        f"Связано {linked_rows} из {total_rows} позиций. "
        f"Не привязано {not_linked_rows} позиций в {OUTPUT_FILE_NOT_ID}."
    )

    return linked_rows, total_rows